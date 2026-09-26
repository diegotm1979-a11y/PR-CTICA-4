#!/usr/bin/env python3
"""Extracts the scouting database from 'Archivo Individualidades.xlsx'
(an export of the .numbers source) into webapp/data.json + webapp/images/.
"""
import json
import re
import sys
import unicodedata
import zipfile
from pathlib import Path

import openpyxl

XLSX_PATH = sys.argv[1] if len(sys.argv) > 1 else "Archivo Individualidades.xlsx"
OUT_DIR = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(".")
IMAGES_DIR = OUT_DIR / "images"


def norm(s):
    if s is None:
        return ""
    return str(s).strip()


def norm_key(s):
    s = norm(s).upper()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", s).strip()


GENERAL_POS_MAP = [
    (("PORTERO",), "Portero"),
    (("CENTRAL", "LATERAL", "CARRILERO"), "Defensa"),
    (("PIVOTE", "INTERIOR", "MEDIO CENTRO", "MEDIA PUNTA"), "Centrocampista"),
    (("BANDA", "DELANTERO", "EXTREMO"), "Delantero"),
]


def posicion_general(posicion):
    if not posicion:
        return "Sin clasificar"
    first = norm(posicion).split("/")[0].strip().upper()
    for keywords, label in GENERAL_POS_MAP:
        if any(k in first for k in keywords):
            return label
    return "Otro"


def slugify(s):
    s = norm_key(s).lower()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s or "jugador"


def col_text(row):
    return [norm(v) for v in row if norm(v)]


def extract_datos_jugador(ws):
    rows = list(ws.iter_rows(min_row=1, max_row=5, values_only=True))
    header = None
    data = None
    for r in rows:
        vals = [norm(v) for v in r]
        if "NOMBRE" in vals:
            header = vals
        elif header is not None and data is None:
            data = [norm(v) for v in r]
    if not header or not data:
        return {}
    out = {}
    for h, v in zip(header, data):
        if h:
            out[h] = v
    return out


def extract_paragraph(ws):
    lines = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        for v in row:
            if norm(v):
                lines.append(norm(v))
    return " ".join(lines).strip()


def extract_list(ws):
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        text = col_text(row)
        if text:
            items.append(" | ".join(text))
    return items


class ImageExtractor:
    """Resolves the picture embedded in a worksheet (e.g. a 'Dibujos' tab)
    by walking the raw OOXML relationships, since openpyxl does not surface
    images exported by Numbers."""

    def __init__(self, xlsx_path):
        self.zip = zipfile.ZipFile(xlsx_path)
        wb_xml = self.zip.read("xl/workbook.xml").decode("utf-8", "ignore")
        self.sheet_rid = dict(
            re.findall(r'<sheet name="([^"]+)"[^>]*r:id="(rId\d+)"', wb_xml)
        )
        rels_xml = self.zip.read("xl/_rels/workbook.xml.rels").decode("utf-8", "ignore")
        rid_target = dict(
            re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels_xml)
        )
        self.sheet_path = {}
        for name, rid in self.sheet_rid.items():
            target = rid_target.get(rid)
            if target:
                self.sheet_path[name] = "xl/" + target.lstrip("/")

    def get_image_bytes(self, sheet_name):
        path = self.sheet_path.get(sheet_name)
        if not path:
            return None
        sheet_file = path.rsplit("/", 1)[-1]
        rels_path = f"xl/worksheets/_rels/{sheet_file}.rels"
        try:
            rels_xml = self.zip.read(rels_path).decode("utf-8", "ignore")
        except KeyError:
            return None
        m = re.search(r'Target="\.\./drawings/([^"]+)"', rels_xml)
        if not m:
            return None
        drawing_file = m.group(1)
        drawing_xml = self.zip.read(f"xl/drawings/{drawing_file}").decode("utf-8", "ignore")
        m2 = re.search(r'r:embed="(rId\d+)"', drawing_xml)
        if not m2:
            return None
        embed_rid = m2.group(1)
        drawing_rels = self.zip.read(
            f"xl/drawings/_rels/{drawing_file}.rels"
        ).decode("utf-8", "ignore")
        m3 = re.search(
            rf'Id="{embed_rid}"[^>]*Target="\.\./media/([^"]+)"', drawing_rels
        )
        if not m3:
            return None
        media_file = m3.group(1)
        return self.zip.read(f"xl/media/{media_file}"), media_file


def main():
    print(f"Loading {XLSX_PATH} ...")
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)

    ws_master = wb["BASE DE DATOS"]
    master = {}
    for row in ws_master.iter_rows(min_row=8, values_only=True):
        jugador = row[1] if len(row) > 1 else None
        if not norm(jugador):
            continue
        key = norm_key(jugador)
        master[key] = {
            "competicion": norm(row[2]) if len(row) > 2 else "",
            "posicion_master": norm(row[3]) if len(row) > 3 else "",
            "equipo": norm(row[4]) if len(row) > 4 else "",
            "edad": row[5] if len(row) > 5 else None,
            "pie_master": norm(row[6]) if len(row) > 6 else "",
        }
    print(f"Master table rows: {len(master)}")

    sheetnames = wb.sheetnames
    groups = []
    cur = None
    for n in sheetnames:
        if n.endswith("Tabla 1"):
            if cur:
                groups.append(cur)
            cur = [n]
        elif cur is not None:
            cur.append(n)
    if cur:
        groups.append(cur)
    groups = [g for g in groups if not g[0].startswith("PLANTILLA")]
    print(f"Detail sheet groups: {len(groups)}")

    img_extractor = ImageExtractor(XLSX_PATH)
    IMAGES_DIR.mkdir(parents=True, exist_ok=True)

    players_by_key = {}

    for g in groups:
        prefix = g[0][: -len(" - Tabla 1")].strip()
        sheet_by_role = {}
        for s in g[1:]:
            suffix = s[len(prefix):].strip(" -")
            sheet_by_role[suffix] = s

        datos_sheet = next((s for role, s in sheet_by_role.items() if role.startswith("DATOS DEL JUGADOR")), None)
        perfil_sheet = next((s for role, s in sheet_by_role.items() if role.startswith("DEFINICI")), None)
        def_sheet = next((s for role, s in sheet_by_role.items() if role.startswith("ASPECTOS DEFENSIV")), None)
        of_sheet = next((s for role, s in sheet_by_role.items() if role.startswith("ASPECTOS OFENSIV")), None)
        notas_sheet = next((s for role, s in sheet_by_role.items() if role.startswith("NOTAS")), None)
        dibujos_sheet = next((s for role, s in sheet_by_role.items() if role.startswith("Dibujos")), None)

        datos = extract_datos_jugador(wb[datos_sheet]) if datos_sheet else {}
        nombre = datos.get("NOMBRE") or prefix
        key = norm_key(nombre) or norm_key(prefix)

        perfil = extract_paragraph(wb[perfil_sheet]) if perfil_sheet else ""
        aspectos_def = extract_list(wb[def_sheet]) if def_sheet else []
        aspectos_of = extract_list(wb[of_sheet]) if of_sheet else []
        notas = extract_list(wb[notas_sheet]) if notas_sheet else []

        image_rel_path = None
        if dibujos_sheet:
            try:
                result = img_extractor.get_image_bytes(dibujos_sheet)
                if result:
                    img_bytes, media_name = result
                    ext = media_name.rsplit(".", 1)[-1].lower()
                    fname = f"{slugify(nombre)}.{ext}"
                    (IMAGES_DIR / fname).write_bytes(img_bytes)
                    image_rel_path = f"images/{fname}"
            except Exception as e:
                print(f"  [img warn] {nombre}: {e}")

        m = master.get(key, {})
        pos = datos.get("POSICIÓN") or m.get("posicion_master") or ""
        players_by_key[key] = {
            "id": slugify(nombre),
            "nombre": nombre,
            "posicion": pos,
            "posicion_general": posicion_general(pos),
            "pie_dominante": datos.get("PIE DOMINANTE") or m.get("pie_master") or "",
            "dorsal": datos.get("DORSAL") or "",
            "altura": datos.get("ALTURA") or "",
            "peso": datos.get("PESO") or "",
            "competicion": m.get("competicion", ""),
            "equipo": m.get("equipo", ""),
            "edad": m.get("edad"),
            "perfil": perfil,
            "aspectos_defensivos": aspectos_def,
            "aspectos_ofensivos": aspectos_of,
            "notas_partidos": notas,
            "imagen": image_rel_path,
            "ficha_completa": True,
        }

    # Add master-only rows (no detail sheet group)
    added_master_only = 0
    for key, m in master.items():
        if key in players_by_key:
            continue
        # reconstruct display name from master (we only stored normalized key there,
        # so re-scan master sheet for original casing)
        added_master_only += 1
        players_by_key[key] = {
            "id": slugify(key),
            "nombre": key.title(),
            "posicion": m.get("posicion_master", ""),
            "posicion_general": posicion_general(m.get("posicion_master", "")),
            "pie_dominante": m.get("pie_master", ""),
            "dorsal": "",
            "altura": "",
            "peso": "",
            "competicion": m.get("competicion", ""),
            "equipo": m.get("equipo", ""),
            "edad": m.get("edad"),
            "perfil": "",
            "aspectos_defensivos": [],
            "aspectos_ofensivos": [],
            "notas_partidos": [],
            "imagen": None,
            "ficha_completa": False,
        }

    # fix master-only display names using original casing from BASE DE DATOS
    orig_names = {}
    for row in ws_master.iter_rows(min_row=8, values_only=True):
        jugador = row[1] if len(row) > 1 else None
        if norm(jugador):
            orig_names[norm_key(jugador)] = norm(jugador)
    for key, p in players_by_key.items():
        if not p["ficha_completa"] and key in orig_names:
            p["nombre"] = orig_names[key]

    players = sorted(players_by_key.values(), key=lambda p: p["nombre"])

    print(f"Total players: {len(players)} (master-only: {added_master_only})")

    competiciones = sorted({p["competicion"] for p in players if p["competicion"]})
    posiciones = sorted({p["posicion"] for p in players if p["posicion"]})
    posiciones_generales = sorted({p["posicion_general"] for p in players})
    equipos = sorted({p["equipo"] for p in players if p["equipo"]})
    pies = sorted({p["pie_dominante"] for p in players if p["pie_dominante"]})
    edades = [p["edad"] for p in players if isinstance(p["edad"], (int, float))]

    import time
    payload = {
        "generated_at": time.time(),
        "count": len(players),
        "filters": {
            "competiciones": competiciones,
            "posiciones": posiciones,
            "posiciones_generales": posiciones_generales,
            "equipos": equipos,
            "pies": pies,
            "edad_min": min(edades) if edades else 15,
            "edad_max": max(edades) if edades else 45,
        },
        "players": players,
    }

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_file = OUT_DIR / "data.json"
    tmp_file = OUT_DIR / "data.json.tmp"
    tmp_file.write_text(json.dumps(payload, ensure_ascii=False, indent=None), encoding="utf-8")
    tmp_file.replace(out_file)
    print(f"Wrote {out_file} ({out_file.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
